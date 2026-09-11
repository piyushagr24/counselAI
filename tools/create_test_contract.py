from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


output = Path("demo_contracts/Legal_Test_Service_Agreement.docx")
output.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(20)
styles["Title"].font.bold = True
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(13)
styles["Heading 1"].font.bold = True

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Service Agreement")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Test contract for AI Legal Contract Assistant")
run.italic = True
run.font.size = Pt(10)

doc.add_paragraph(
    "This Service Agreement is entered into on January 15, 2026, by and between Northwind Analytics LLC (the Vendor) and Acme Retail Inc. (the Client). The parties agree to the following terms."
)

sections = [
    ("1 Services", "The Vendor will provide hosted sales analytics software, account support, and monthly performance reports. The Client may use the software for its internal business operations during the Term."),
    ("2 Payment Terms", "The Client will pay a subscription fee of $1,200 per month. Invoices are due within 15 days after receipt. Late amounts accrue interest at 1% per month, or the maximum amount permitted by law, whichever is lower."),
    ("3 Term and Renewal", "The initial term begins January 15, 2026 and ends January 14, 2027. The Agreement automatically renews for additional one-year periods unless either party gives at least 30 days written notice of non-renewal."),
    ("4 Confidentiality", "Each party must protect the other party's confidential information using reasonable care and may disclose it only to personnel who need the information to perform this Agreement. These obligations continue for three years after termination."),
    ("5 Intellectual Property", "The Vendor retains all rights in the software, documentation, and improvements. The Client retains all rights in data it submits to the software and grants the Vendor a limited license to process that data to provide the services."),
    ("6 Limitation of Liability", "Neither party is liable for indirect, incidental, special, or consequential damages. The Vendor's total aggregate liability under this Agreement will not exceed the fees paid by the Client during the six months before the event giving rise to the claim."),
    ("7 Termination", "Either party may terminate this Agreement for a material breach if the breach is not cured within 30 days after written notice. The Client must pay all undisputed fees accrued through the termination date."),
    ("8 Governing Law", "This Agreement is governed by the laws of the State of Delaware, without regard to its conflict-of-law rules. The parties consent to the exclusive jurisdiction of the state and federal courts located in Delaware."),
]

for heading, body in sections:
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body)

doc.add_heading("Signatures", level=1)
table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
table.cell(0, 0).text = "Northwind Analytics LLC"
table.cell(0, 1).text = "Acme Retail Inc."
table.cell(1, 0).text = "By: ______________________________"
table.cell(1, 1).text = "By: ______________________________"
table.cell(2, 0).text = "Date: ____________________________"
table.cell(2, 1).text = "Date: ____________________________"

doc.save(output)
print(output.resolve())
