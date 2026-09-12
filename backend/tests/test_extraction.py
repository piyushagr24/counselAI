import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pymupdf
import docx
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def _auth_headers(email: str = "tester_extract@example.com", password: str = "password123") -> dict:
    signup_res = client.post(
        "/api/auth/signup",
        json={"email": email, "password": password, "name": "Extract Tester"},
    )
    if signup_res.status_code == 200:
        tok = signup_res.json()["access_token"]
    else:
        login_res = client.post(
            "/api/auth/login",
            json={"email": email, "password": password},
        )
        tok = login_res.json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _make_sample_pdf_bytes() -> bytes:
    pdf = pymupdf.open()
    page1 = pdf.new_page()
    page1.insert_text((72, 72), "PAGE ONE: This Agreement is entered into by Party A and Party B.")
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "PAGE TWO: Termination clause. Either party may terminate with 30 days notice.")
    buf = io.BytesIO(pdf.tobytes())
    pdf.close()
    return buf.getvalue()


def _make_sample_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Confidentiality", level=1)
    document.add_paragraph("The Receiving Party shall keep all information confidential.")
    document.add_heading("Payment Terms", level=1)
    document.add_paragraph("Payment is due within 15 days of invoice.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_pdf_extraction_preserves_pages():
    headers = _auth_headers()
    pdf_bytes = _make_sample_pdf_bytes()
    response = client.post(
        "/api/contracts/upload",
        files={"file": ("sample.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
        headers=headers,
    )
    assert response.status_code == 200
    body = response.json()
    extraction = body["extraction"]

    assert extraction["metadata"]["source_type"] == "pdf"
    assert extraction["metadata"]["num_pages"] == 2
    assert len(extraction["segments"]) == 2
    assert extraction["segments"][0]["page_number"] == 1
    assert extraction["segments"][1]["page_number"] == 2
    assert "Termination clause" in extraction["full_text"]


def test_docx_extraction_captures_headings():
    headers = _auth_headers()
    docx_bytes = _make_sample_docx_bytes()
    response = client.post(
        "/api/contracts/upload",
        files={
            "file": (
                "sample.docx",
                io.BytesIO(docx_bytes),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=headers,
    )
    assert response.status_code == 200
    body = response.json()
    extraction = body["extraction"]

    assert extraction["metadata"]["source_type"] == "docx"
    assert extraction["metadata"]["num_pages"] == 1
    contract_res = client.get(f"/api/contracts/{body['contract_id']}", headers=headers)
    assert contract_res.status_code == 200
    assert contract_res.json()["num_pages"] == 1
    headings = [s["heading"] for s in extraction["segments"] if s["heading"]]
    assert any("Heading" in h for h in headings)
    assert "Payment is due within 15 days" in extraction["full_text"]
