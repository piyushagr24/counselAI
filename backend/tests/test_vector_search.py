import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import docx
from fastapi.testclient import TestClient

from app.main import app
from app.services.vector_store import query_chunks

client = TestClient(app)


def _make_sample_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Confidentiality", level=1)
    document.add_paragraph(
        "The Receiving Party shall keep all confidential information secret "
        "and shall not disclose it to any third party without prior written consent."
    )
    document.add_heading("Payment Terms", level=1)
    document.add_paragraph(
        "The Client shall pay the Contractor within 15 days of receiving each invoice."
    )
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _auth_headers(email: str = "vector_tester@example.com", password: str = "password123") -> dict:
    signup_res = client.post(
        "/api/auth/signup",
        json={"email": email, "password": password, "name": "Vector Tester"},
    )
    if signup_res.status_code == 200:
        tok = signup_res.json()["access_token"]
    else:
        login_res = client.post(
            "/api/auth/login",
            json={"email": email, "password": password},
        )
        tok = login_res.json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def test_similarity_search_retrieves_relevant_chunk():
    headers = _auth_headers()
    response = client.post(
        "/api/contracts/upload",
        headers=headers,
        files={
            "file": (
                "sample.docx",
                io.BytesIO(_make_sample_docx_bytes()),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    body = response.json()
    contract_id = body["contract_id"]
    assert body["chunks_indexed"] == 2

    results = query_chunks(
        "Can the receiving party share confidential information with others?",
        contract_id=contract_id,
        top_k=1,
    )

    assert len(results) == 1
    assert "confidential" in results[0]["text"].lower()
    assert results[0]["heading"] == "Heading 1"
    assert results[0]["contract_id"] == contract_id
